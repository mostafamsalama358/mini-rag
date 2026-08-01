"""Re-index project corpus from RAG_CORPUS workbook (PRODUCTS + INTERACTIONS + ALIASES)."""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import time
from collections import defaultdict
from pathlib import Path

import openpyxl
import requests

from _render_workbook_product import (
    alias_file_name,
    interaction_file_name,
    render_alias_document,
    render_interaction_document,
    render_product,
)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "excel sheets" / "RAG_CORPUS_56.product_patched.xlsx"
DEFAULT_OUT = ROOT / "eval_run" / "e2e_live_results"
DEFAULT_RENDER_DIR = ROOT / "eval_run" / "workbook_corpus"


def _load_workbook(
    path: Path,
) -> tuple[list[dict], list[dict], dict[str, list[dict]], dict[str, list[dict]]]:
    if path.name.startswith("~$"):
        raise SystemExit(
            f"Refusing Excel lock file: {path.name}\n"
            f"Use the real workbook: RAG_CORPUS_56.product_patched.xlsx"
        )
    if not path.exists():
        raise SystemExit(f"Workbook not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws_products = wb["PRODUCTS"]
    headers = [cell.value for cell in ws_products[1]]
    products = [dict(zip(headers, row)) for row in ws_products.iter_rows(min_row=2, values_only=True)]

    interactions: list[dict] = []
    interactions_by_sku: dict[str, list[dict]] = defaultdict(list)
    if "INTERACTIONS" in wb.sheetnames:
        ws_inter = wb["INTERACTIONS"]
        inter_headers = [cell.value for cell in ws_inter[1]]
        for row in ws_inter.iter_rows(min_row=2, values_only=True):
            item = dict(zip(inter_headers, row))
            sku_id = str(item.get("sku_id") or "").strip()
            if not sku_id:
                continue
            interactions.append(item)
            interactions_by_sku[sku_id].append(item)

    aliases_by_sku: dict[str, list[dict]] = defaultdict(list)
    if "ALIASES" in wb.sheetnames:
        ws_aliases = wb["ALIASES"]
        alias_headers = [cell.value for cell in ws_aliases[1]]
        for row in ws_aliases.iter_rows(min_row=2, values_only=True):
            item = dict(zip(alias_headers, row))
            sku_id = str(item.get("sku_id") or "").strip()
            if sku_id:
                aliases_by_sku[sku_id].append(item)

    return products, interactions, dict(interactions_by_sku), dict(aliases_by_sku)


def _write(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def _render_documents(
    products: list[dict],
    interactions: list[dict],
    interactions_by_sku: dict[str, list[dict]],
    aliases_by_sku: dict[str, list[dict]],
    render_dir: Path,
) -> tuple[list[Path], dict[str, int]]:
    if render_dir.exists():
        shutil.rmtree(render_dir)
    render_dir.mkdir(parents=True, exist_ok=True)

    written: list[Path] = []
    counts = {"products": 0, "interactions": 0, "aliases": 0}

    for product in products:
        file_name = str(product.get("file_name") or "").strip()
        if not file_name:
            continue
        sku_id = str(product.get("sku_id") or "").strip()
        text = render_product(product, interactions_by_sku.get(sku_id, []))
        written.append(_write(render_dir / file_name, text))
        counts["products"] += 1

    seen_interaction_names: set[str] = set()
    for row in interactions:
        rel_name = interaction_file_name(row)
        if rel_name in seen_interaction_names:
            continue
        seen_interaction_names.add(rel_name)
        written.append(_write(render_dir / rel_name, render_interaction_document(row)))
        counts["interactions"] += 1

    for product in products:
        sku_id = str(product.get("sku_id") or "").strip()
        brand = str(product.get("brand_name") or "").strip()
        aliases = aliases_by_sku.get(sku_id, [])
        if not brand or not aliases:
            continue
        rel_name = alias_file_name(brand, sku_id)
        written.append(_write(render_dir / rel_name, render_alias_document(brand, sku_id, aliases)))
        counts["aliases"] += 1

    return sorted(written), counts


def clear_project_corpus(project_id: int) -> None:
    """Remove old assets/files/chunks for a workbook-only re-index."""
    print(f"Clearing project {project_id} assets, chunks, vectors, and files…")
    collection = f"collection_768_{project_id}"
    sql = (
        f"TRUNCATE TABLE {collection}; "
        f"DELETE FROM chunks WHERE chunk_project_id={project_id}; "
        f"DELETE FROM assets WHERE asset_project_id={project_id};"
    )
    subprocess.run(
        [
            "docker",
            "exec",
            "pgvector",
            "psql",
            "-U",
            "postgres",
            "-d",
            "algorag",
            "-c",
            sql,
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    subprocess.run(
        [
            "docker",
            "exec",
            "fastapi",
            "sh",
            "-c",
            f"rm -rf /app/assets/files/{project_id}/*",
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    print(f"Cleared project {project_id} (vectors, DB assets/chunks, files directory)")


def wait_task(api: str, task_id: str, timeout: float = 1800.0) -> dict:
    deadline = time.time() + timeout
    last: dict = {}
    while time.time() < deadline:
        try:
            response = requests.get(f"{api}/api/v1/data/tasks/{task_id}", timeout=60)
            if response.status_code >= 500:
                print(f"  task={task_id[:8]}… poll HTTP {response.status_code}; retrying")
                time.sleep(10)
                continue
            response.raise_for_status()
            last = response.json()
        except requests.RequestException as exc:
            print(f"  task={task_id[:8]}… poll error: {exc}; retrying")
            time.sleep(10)
            continue
        state = last.get("status") or last.get("state")
        ready = last.get("ready")
        print(f"  task={task_id[:8]}… ready={ready} status={state}")
        if ready:
            return last
        time.sleep(5)
    raise TimeoutError(f"task timeout: {last}")


def upload_file(api: str, project_id: int, path: Path, render_dir: Path) -> str:
    upload_name = str(path.relative_to(render_dir)).replace("\\", "/").replace("/", "__")
    with path.open("rb") as handle:
        response = requests.post(
            f"{api}/api/v1/data/upload/{project_id}",
            files={"file": (upload_name, handle, "text/plain")},
            timeout=120,
        )
    if response.status_code >= 400:
        raise RuntimeError(f"upload {path.name}: {response.status_code} {response.text}")
    body = response.json()
    if body.get("signal") != "file_upload_success":
        raise RuntimeError(f"upload signal: {body}")
    return body.get("asset_name") or upload_name


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Workbook-only re-index (3 sheets)")
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Path to RAG_CORPUS workbook (default: product_patched.xlsx)",
    )
    parser.add_argument("--project-id", type=int, default=2)
    parser.add_argument("--api", default="http://localhost:8000")
    parser.add_argument(
        "--render-dir",
        type=Path,
        default=DEFAULT_RENDER_DIR,
        help="Where rendered corpus files are written before upload",
    )
    parser.add_argument(
        "--render-only",
        action="store_true",
        help="Only render workbook sheets to text files; skip cleanup/upload/re-index",
    )
    parser.add_argument(
        "--skip-cleanup",
        action="store_true",
        help="Do not delete existing project assets before upload",
    )
    args = parser.parse_args()

    products, interactions, interactions_by_sku, aliases_by_sku = _load_workbook(args.workbook)
    files, counts = _render_documents(
        products,
        interactions,
        interactions_by_sku,
        aliases_by_sku,
        args.render_dir,
    )
    print(
        f"Rendered from {args.workbook.name}: "
        f"products={counts['products']} interactions={counts['interactions']} "
        f"aliases={counts['aliases']} total_files={len(files)} -> {args.render_dir}"
    )

    if args.render_only:
        return

    if not args.skip_cleanup:
        clear_project_corpus(args.project_id)

    print(f"Uploading {len(files)} workbook documents to project {args.project_id}…")
    uploaded: list[str] = []
    for path in files:
        asset_name = upload_file(args.api, args.project_id, path, args.render_dir)
        uploaded.append(asset_name)
        print(f"  OK {path.relative_to(args.render_dir)} -> {asset_name}")

    print("process-and-push do_reset=1…")
    response = requests.post(
        f"{args.api}/api/v1/data/process-and-push/{args.project_id}",
        json={"do_reset": 1},
        timeout=60,
    )
    response.raise_for_status()
    body = response.json()
    print(json.dumps(body, indent=2)[:800])
    task_id = body.get("task_id") or body.get("job_id")
    if not task_id:
        raise RuntimeError(f"no task_id in {body}")

    result = wait_task(args.api, str(task_id))
    print("task result:", json.dumps(result, indent=2)[:1200])

    info = requests.get(f"{args.api}/api/v1/nlp/index/info/{args.project_id}", timeout=60)
    info.raise_for_status()
    info_body = info.json()
    print("index info:", json.dumps(info_body, indent=2)[:1500])

    DEFAULT_OUT.mkdir(parents=True, exist_ok=True)
    stamp = {
        "timestamp": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "workbook": str(args.workbook),
        "render_dir": str(args.render_dir),
        "render_counts": counts,
        "uploaded_count": len(uploaded),
        "uploaded": uploaded,
        "process_task": result,
        "index_info": info_body,
    }
    out_path = DEFAULT_OUT / "reindex_workbook_3sheets.json"
    out_path.write_text(json.dumps(stamp, indent=2), encoding="utf-8")
    print(f"wrote {out_path}")


if __name__ == "__main__":
    main()
