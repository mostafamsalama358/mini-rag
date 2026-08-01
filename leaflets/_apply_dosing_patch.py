"""One-off patch: apply curated strength/min_age/adult_dose to RAG_CORPUS_56.xlsx."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

DOSING_BY_FILE: dict[str, dict[str, str]] = {
    p["file_name"]: {
        "strength": p["strength"],
        "min_age": p["min_age"],
        "adult_dose": p["adult_dose"],
    }
    for p in [
    {
        "file_name": "01_panadol_paracetamol_500.txt",
        "strength": "Paracetamol 500 mg/tablet",
        "min_age": "6 years",
        "adult_dose": "1-2 tablets every 4-6 hours; max 8 tablets/day (4 g)",
    },
    {
        "file_name": "02_panadol_cold_flu_day.txt",
        "strength": "Paracetamol 500 mg + Phenylephrine HCl 5 mg/tablet",
        "min_age": "12 years",
        "adult_dose": "2 tablets every 4-6 hours; max 8 tablets/day",
    },
    {
        "file_name": "03_panadol_sinus_relief_pe.txt",
        "strength": "Paracetamol 500 mg + Phenylephrine HCl 5 mg/tablet",
        "min_age": "12 years",
        "adult_dose": "2 tablets every 4-6 hours; max 8 tablets/day",
    },
    {
        "file_name": "04_cetal_cold_flu_day.txt",
        "strength": "Paracetamol 500 mg + Phenylephrine HCl 5 mg + Caffeine 25 mg",
        "min_age": "12 years",
        "adult_dose": "2 caplets every 4-6 hours; max 8/day",
    },
    {
        "file_name": "05_flurest_n.txt",
        "strength": "Paracetamol 450 mg + Pseudoephedrine 30 mg + Chlorpheniramine 2 mg",
        "min_age": "12 years",
        "adult_dose": "1 tablet 3-4 times/day",
    },
    {
        "file_name": "06_congestal.txt",
        "strength": "Paracetamol 450 mg + Pseudoephedrine 30 mg + Chlorpheniramine 2 mg",
        "min_age": "12 years",
        "adult_dose": "1 tablet 3 times/day",
    },
    {
        "file_name": "07_sine_up.txt",
        "strength": "Pseudoephedrine + Triprolidine (combination)",
        "min_age": "12 years",
        "adult_dose": "1 tablet every 6-8 hours",
    },
    {
        "file_name": "08_bradozen.txt",
        "strength": "Herbal cough formulation",
        "min_age": "12 years",
        "adult_dose": "1 tablet 3 times/day",
    },
    {
        "file_name": "09_notussil.txt",
        "strength": "Levodropropizine syrup",
        "min_age": "2 years",
        "adult_dose": "10 mL three times/day",
    },
    {
        "file_name": "10_tussigreen.txt",
        "strength": "Herbal cough syrup",
        "min_age": "2 years",
        "adult_dose": "10 mL three times/day",
    },
    {
        "file_name": "11_daflon_500.txt",
        "strength": "Diosmin 450 mg + Hesperidin 50 mg",
        "min_age": "18 years",
        "adult_dose": "2 tablets/day",
    },
    {
        "file_name": "12_daflon_1000.txt",
        "strength": "Diosmin 900 mg + Hesperidin 100 mg",
        "min_age": "18 years",
        "adult_dose": "1 tablet/day",
    },
    {
        "file_name": "13_controloc_20.txt",
        "strength": "Pantoprazole 20 mg",
        "min_age": "12 years",
        "adult_dose": "20 mg once daily",
    },
    {
        "file_name": "14_jeparilon.txt",
        "strength": "Heparinoid topical cream",
        "min_age": "12 years",
        "adult_dose": "Apply 2-3 times/day",
    },
    {
        "file_name": "15_anaseziago.txt",
        "strength": "Topical anesthetic gel",
        "min_age": "12 years",
        "adult_dose": "Apply as needed up to 3-4 times/day",
    },
    {
        "file_name": "16_adol_500.txt",
        "strength": "Paracetamol 500 mg",
        "min_age": "6 years",
        "adult_dose": "1-2 tablets every 4-6 hours; max 4 g/day",
    },
    {
        "file_name": "17_cataflam_25.txt",
        "strength": "Diclofenac potassium 25 mg",
        "min_age": "14 years",
        "adult_dose": "1 tablet 2-3 times/day",
    },
    {
        "file_name": "18_brufen_400.txt",
        "strength": "Ibuprofen 400 mg",
        "min_age": "12 years",
        "adult_dose": "1 tablet every 6-8 hours; max 1200-2400 mg/day",
    },
    {
        "file_name": "19_antinal.txt",
        "strength": "Nifuroxazide 200 mg",
        "min_age": "6 years",
        "adult_dose": "1 capsule 4 times/day",
    },
    {
        "file_name": "20_flagyl_500.txt",
        "strength": "Metronidazole 500 mg",
        "min_age": "10 years",
        "adult_dose": "500 mg every 8-12 hours",
    },
    {
        "file_name": "21_augmentin_1g.txt",
        "strength": "Amoxicillin 875 mg + Clavulanic acid 125 mg",
        "min_age": "12 years or ≥40 kg",
        "adult_dose": "1 tablet every 12 hours",
    },
    {
        "file_name": "22_hibiotic_1g.txt",
        "strength": "Amoxicillin 875 mg + Clavulanic acid 125 mg",
        "min_age": "12 years or ≥40 kg",
        "adult_dose": "1 tablet every 12 hours",
    },
    {
        "file_name": "23_zithrokan_500.txt",
        "strength": "Azithromycin 500 mg",
        "min_age": "45 kg or adult",
        "adult_dose": "500 mg once daily for 3 days",
    },
    {
        "file_name": "24_ciprobay_500.txt",
        "strength": "Ciprofloxacin 500 mg",
        "min_age": "18 years",
        "adult_dose": "500 mg every 12 hours",
    },
    {
        "file_name": "25_risek_20.txt",
        "strength": "Omeprazole 20 mg",
        "min_age": "12 years",
        "adult_dose": "20 mg once daily",
    },
    {
        "file_name": "26_gaviscon.txt",
        "strength": "Alginate + Antacid suspension/tablets",
        "min_age": "12 years",
        "adult_dose": "10-20 mL or 2-4 tablets after meals and bedtime",
    },
    {
        "file_name": "27_motilium_10.txt",
        "strength": "Domperidone 10 mg",
        "min_age": "12 years and ≥35 kg",
        "adult_dose": "10 mg up to 3 times/day before meals",
    },
    {
        "file_name": "28_buscopan.txt",
        "strength": "Hyoscine butylbromide 10 mg",
        "min_age": "6 years",
        "adult_dose": "1-2 tablets 3-5 times/day",
    },
    {
        "file_name": "29_concor_5.txt",
        "strength": "Bisoprolol 5 mg",
        "min_age": "18 years",
        "adult_dose": "5 mg once daily",
    },
    {
        "file_name": "30_norvasc_5.txt",
        "strength": "Amlodipine 5 mg",
        "min_age": "6 years",
        "adult_dose": "5 mg once daily",
    },
    {
        "file_name": "31_glucophage_500.txt",
        "strength": "Metformin 500 mg",
        "min_age": "10 years",
        "adult_dose": "500 mg 2-3 times/day with meals",
    },
    {
        "file_name": "32_amaryl_2.txt",
        "strength": "Glimepiride 2 mg",
        "min_age": "18 years",
        "adult_dose": "1-2 mg once daily initially",
    },
    {
        "file_name": "33_lipitor_20.txt",
        "strength": "Atorvastatin 20 mg",
        "min_age": "10 years",
        "adult_dose": "20 mg once daily",
    },
    {
        "file_name": "34_plavix_75.txt",
        "strength": "Clopidogrel 75 mg",
        "min_age": "18 years",
        "adult_dose": "75 mg once daily",
    },
    {
        "file_name": "35_aspocid_75.txt",
        "strength": "Aspirin 75 mg",
        "min_age": "16 years",
        "adult_dose": "75 mg once daily",
    },
    {
        "file_name": "36_telfast_120.txt",
        "strength": "Fexofenadine 120 mg",
        "min_age": "12 years",
        "adult_dose": "120 mg once daily",
    },
    {
        "file_name": "37_claritine_10.txt",
        "strength": "Loratadine 10 mg",
        "min_age": "2 years",
        "adult_dose": "10 mg once daily",
    },
    {
        "file_name": "38_otrivin_0_1.txt",
        "strength": "Xylometazoline 0.1%",
        "min_age": "12 years",
        "adult_dose": "1 spray each nostril up to 3 times/day",
    },
    {
        "file_name": "39_ventolin_inhaler.txt",
        "strength": "Salbutamol 100 mcg/puff",
        "min_age": "4 years",
        "adult_dose": "1-2 puffs as needed",
    },
    {
        "file_name": "40_singulair_10.txt",
        "strength": "Montelukast 10 mg",
        "min_age": "15 years",
        "adult_dose": "10 mg once daily evening",
    },
    {
        "file_name": "41_diflucan_150.txt",
        "strength": "Fluconazole 150 mg",
        "min_age": "18 years",
        "adult_dose": "150 mg single dose",
    },
    {
        "file_name": "42_lamisil_250.txt",
        "strength": "Terbinafine 250 mg",
        "min_age": "18 years",
        "adult_dose": "250 mg once daily",
    },
    {
        "file_name": "43_duphaston_10.txt",
        "strength": "Dydrogesterone 10 mg",
        "min_age": "18 years",
        "adult_dose": "Depends indication; commonly 10 mg 1-2 times/day",
    },
    {
        "file_name": "44_folic_acid_5.txt",
        "strength": "Folic acid 5 mg",
        "min_age": "Any age when indicated",
        "adult_dose": "5 mg once daily",
    },
    {
        "file_name": "45_ferrotron.txt",
        "strength": "Iron supplement",
        "min_age": "12 years",
        "adult_dose": "1 capsule once daily",
    },
    {
        "file_name": "46_milga.txt",
        "strength": "Benfotiamine + Vitamin B6 + Vitamin B12",
        "min_age": "18 years",
        "adult_dose": "1 tablet 3 times/day",
    },
    {
        "file_name": "47_smecta.txt",
        "strength": "Diosmectite 3 g/sachet",
        "min_age": "Any age",
        "adult_dose": "3 sachets/day",
    },
    {
        "file_name": "48_enterogermina.txt",
        "strength": "Bacillus clausii 2 billion spores/5 mL",
        "min_age": "Any age",
        "adult_dose": "2-3 vials/day",
    },
    {
        "file_name": "49_strepsils.txt",
        "strength": "Amylmetacresol 0.6 mg + Dichlorobenzyl alcohol 1.2 mg",
        "min_age": "6 years",
        "adult_dose": "1 lozenge every 2-3 hours; max 12/day",
    },
    {
        "file_name": "50_comtrex.txt",
        "strength": "Paracetamol + Pseudoephedrine + Dextromethorphan",
        "min_age": "12 years",
        "adult_dose": "2 tablets every 6 hours",
    },
    {
        "file_name": "51_panadol_extra.txt",
        "strength": "Paracetamol 500 mg + Caffeine 65 mg",
        "min_age": "12 years",
        "adult_dose": "1-2 tablets every 4-6 hours; max 8/day",
    },
    {
        "file_name": "52_panadol_joint.txt",
        "strength": "Paracetamol 665 mg modified-release",
        "min_age": "18 years",
        "adult_dose": "2 tablets every 6-8 hours; max 6/day",
    },
    {
        "file_name": "53_panadol_migraine.txt",
        "strength": "Paracetamol 250 mg + Aspirin 250 mg + Caffeine 65 mg",
        "min_age": "18 years",
        "adult_dose": "2 tablets at onset; max 8/day",
    },
    {
        "file_name": "54_panadol_acute_head_cold.txt",
        "strength": "Paracetamol 500 mg + Phenylephrine HCl 5 mg",
        "min_age": "12 years",
        "adult_dose": "2 tablets every 4-6 hours",
    },
    {
        "file_name": "55_panadol_cold_flu_all_in_one.txt",
        "strength": "Paracetamol 500 mg + Guaifenesin 100 mg + Phenylephrine HCl 5 mg",
        "min_age": "16 years",
        "adult_dose": "2 tablets every 4-6 hours; max 8/day",
    },
    {
        "file_name": "56_panadol_vapour_release.txt",
        "strength": "Paracetamol 500 mg + Guaifenesin 100 mg + Phenylephrine HCl 5 mg",
        "min_age": "12 years",
        "adult_dose": "2 tablets every 4-6 hours; max 8/day",
    },
]
}

PATCHES = [
    {"file_name": fn, **fields} for fn, fields in DOSING_BY_FILE.items()
]

XLSX = Path(r"d:\mini-rag\excel sheets\RAG_CORPUS_56.xlsx")


def main() -> None:
    by_file = {p["file_name"]: p for p in PATCHES}
    assert len(by_file) == 56, f"expected 56 patches, got {len(by_file)}"

    wb = openpyxl.load_workbook(XLSX)
    ws = wb["PRODUCTS"]
    headers = [cell.value for cell in ws[1]]
    col = {name: headers.index(name) + 1 for name in ("file_name", "strength", "min_age", "adult_dose")}

    updated = 0
    missing: list[str] = []
    for row_idx in range(2, ws.max_row + 1):
        file_name = ws.cell(row=row_idx, column=col["file_name"]).value
        if file_name not in by_file:
            missing.append(str(file_name))
            continue
        patch = by_file[file_name]
        ws.cell(row=row_idx, column=col["strength"], value=patch["strength"])
        ws.cell(row=row_idx, column=col["min_age"], value=patch["min_age"])
        ws.cell(row=row_idx, column=col["adult_dose"], value=patch["adult_dose"])
        updated += 1

    if missing:
        raise SystemExit(f"Unmatched PRODUCT rows: {missing}")

    try:
        wb.save(XLSX)
        print(f"Updated {updated}/56 rows in {XLSX}")
    except PermissionError:
        fallback = XLSX.with_suffix(".patched.xlsx")
        wb.save(fallback)
        print(f"Original locked — saved {updated}/56 rows to {fallback}")
        print("Close RAG_CORPUS_56.xlsx in Excel, then run this script again to replace the original.")
        return

    # Verify
    wb2 = openpyxl.load_workbook(XLSX, data_only=True)
    ws2 = wb2["PRODUCTS"]
    headers2 = [cell.value for cell in ws2[1]]
    idx = {name: headers2.index(name) for name in ("file_name", "strength", "min_age", "adult_dose")}
    filled = 0
    for row in ws2.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers2, row))
        if all(d.get(k) for k in ("strength", "min_age", "adult_dose")):
            filled += 1
    print(f"Verified filled strength/min_age/adult_dose: {filled}/56")
    sample = next(dict(zip(headers2, r)) for r in ws2.iter_rows(min_row=2, values_only=True) if dict(zip(headers2, r))["file_name"] == "09_notussil.txt")
    print("sample 09_notussil:", json.dumps({k: sample[k] for k in ("strength", "min_age", "adult_dose")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
