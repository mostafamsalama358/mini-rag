"""Apply curated product patch fields to the patched workbook."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

from _apply_product_patch import PRODUCT_PATCH_BY_FILE

sys.stdout.reconfigure(encoding="utf-8")

XLSX = Path(r"d:\mini-rag\excel sheets\RAG_CORPUS_56.patched.xlsx")

ORDERED_COLUMNS = [
    "max_daily_dose",
    "max_daily_dose_value",
    "max_daily_dose_unit",
    "cross_selling",
    "pediatric_dose",
    "pediatric_min_age",
    "pediatric_max_daily_dose",
]


def _headers(ws) -> list[str]:
    return [cell.value for cell in ws[1]]


def _ensure_pediatric_columns(ws) -> None:
    headers = _headers(ws)
    if "pediatric_min_age" in headers and "pediatric_max_daily_dose" in headers:
        return

    ped_idx = headers.index("pediatric_dose") + 1
    ws.insert_cols(ped_idx + 1, amount=2)
    ws.cell(row=1, column=ped_idx + 1, value="pediatric_min_age")
    ws.cell(row=1, column=ped_idx + 2, value="pediatric_max_daily_dose")


def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["PRODUCTS"]
    _ensure_pediatric_columns(ws)
    headers = _headers(ws)
    col = {name: headers.index(name) + 1 for name in ["file_name", *ORDERED_COLUMNS] if name in headers}

    updated = 0
    for row_idx in range(2, ws.max_row + 1):
        file_name = ws.cell(row=row_idx, column=col["file_name"]).value
        patch = PRODUCT_PATCH_BY_FILE.get(file_name)
        if not patch:
            continue
        for key, value in patch.items():
            ws.cell(row=row_idx, column=col[key], value=value)
        updated += 1

    saved_path = XLSX
    try:
        wb.save(XLSX)
    except PermissionError:
        saved_path = XLSX.with_name("RAG_CORPUS_56.product_patched.xlsx")
        wb.save(saved_path)
        print(f"Original workbook locked; saved updated copy to {saved_path.name}")
    else:
        print(f"Updated {updated} rows in {XLSX.name}")

    # Verify counts after patch.
    wb2 = openpyxl.load_workbook(saved_path, data_only=True)
    ws2 = wb2["PRODUCTS"]
    headers2 = _headers(ws2)
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    for field in ORDERED_COLUMNS:
        idx = headers2.index(field)
        empty = 0
        for row in rows:
            value = row[idx]
            if value is None or str(value).strip() in ("", "-", "None", "N/A", "n/a"):
                empty += 1
        print(f"{field}: {empty}/56 empty")


if __name__ == "__main__":
    main()
