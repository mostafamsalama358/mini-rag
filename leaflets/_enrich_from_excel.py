"""Enrich leaflets/ from Egyptian Excel catalog sheets.

Sources:
  - excel sheets/FINAL MEDICIN.xlsx
  - excel sheets/ALL MATERIAL 1.xlsx
  - excel sheets/NEW DRUG INTERACTION.xlsx
  - excel sheets/FINAL MATERIAL VS DISEASE.xlsx
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

ROOT = Path(r"d:\mini-rag")
OUT = ROOT / "leaflets"
EXCEL = ROOT / "excel sheets"

# leaflet_file -> preferred FINAL MEDICIN med substring match (upper)
SKU_PREFERS: dict[str, list[str]] = {
    "01_panadol_paracetamol_500.txt": ["PANADOL ADVANCE 500", "PANADOL ADVANCE 48"],
    "02_panadol_cold_flu_day.txt": ["PANADOL COLD & FLU DAY"],
    "03_panadol_sinus_relief_pe.txt": ["PANADOL SINUS RELIEF PE"],
    "51_panadol_extra.txt": ["PANADOL EXTRA 24", "PANADOL EXTRA 48"],
    "52_panadol_joint.txt": ["PANADOL JOINT 24", "PANADOL JOINT 18"],
    "53_panadol_migraine.txt": ["PANADOL MIGRAINE"],
    "54_panadol_acute_head_cold.txt": ["PANADOL ACUTE HEAD COLD"],
    "55_panadol_cold_flu_all_in_one.txt": ["PANADOL COLD & FLU (ALL IN ONE)"],
    "56_panadol_vapour_release.txt": ["PANADOL COLD & FLU VAPOUR RELEASE"],
    "04_cetal_cold_flu_day.txt": ["CETAL COLD & FLU"],
    "05_flurest_n.txt": ["FLUREST N"],
    "06_congestal.txt": ["CONGESTAL 20 TABS"],
    "07_sine_up.txt": ["SINE UP 20 TAB"],
    "08_bradozen.txt": ["BRADOZEN"],
    "09_notussil.txt": ["NOTUSSIL 4MG"],
    "10_tussigreen.txt": ["TUSSIGLOBE", "TUSSIPECT"],  # closest catalog cough brands
    "11_daflon_500.txt": ["DAFLON 500"],
    "12_daflon_1000.txt": ["DAFLON 1000"],
    "13_controloc_20.txt": ["CONTROLOC 20MG"],
    "14_jeparilon.txt": ["JEPARILON"],
    "15_anaseziago.txt": ["ANASEZIAGO"],
    "16_adol_500.txt": ["ADOL 500MG"],
    "17_cataflam_25.txt": ["CATAFLAM 50 MG", "CATAFLAM 25 MG"],
    "18_brufen_400.txt": ["BRUFEN 400 MG"],
    "19_antinal.txt": ["ANTINAL 200MG 24"],
    "20_flagyl_500.txt": ["FLAGYL 500MG 20 TAB", "FLAGYL 500MG 20", "FLAGYL 250MG 20 TAB"],
    "21_augmentin_1g.txt": ["AUGMENTIN 1 GM"],
    "22_hibiotic_1g.txt": ["HIBIOTIC 1 GM"],
    "23_zithrokan_500.txt": ["ZITHROKAN 500"],
    "24_ciprobay_500.txt": ["CIPROBAY 500"],
    "25_risek_20.txt": ["RISEK 20 MG"],
    "26_gaviscon.txt": ["GAVISCON ADVANCE", "GAVISCON LIQUID"],
    "27_motilium_10.txt": ["MOTILIUM 10MG"],
    "28_buscopan.txt": ["BUSCOPAN PLUS", "BUSCOPAN COMPOSITUM 20 SUGAR"],
    "29_concor_5.txt": ["CONCOR 5 MG"],
    "30_norvasc_5.txt": ["NORVASC 5MG"],
    "31_glucophage_500.txt": ["GLUCOPHAGE 500 MG"],
    "32_amaryl_2.txt": ["AMARYL 2 MG"],
    "33_lipitor_20.txt": ["LIPITOR 20 MG"],
    "34_plavix_75.txt": ["PLAVIX 75"],
    "35_aspocid_75.txt": ["ASPOCID 75MG"],
    "36_telfast_120.txt": ["TELFAST 120"],
    "37_claritine_10.txt": ["CLARITINE 10MG"],
    "38_otrivin_0_1.txt": ["OTRIVIN 0.1%"],
    "39_ventolin_inhaler.txt": ["VENTOLIN 100", "VENTOLIN EVOHALER", "VENTOLIN 0.5% RESPIRATOR"],
    "40_singulair_10.txt": ["SINGULAIR 10MG"],
    "41_diflucan_150.txt": ["DIFLUCAN 150"],
    "42_lamisil_250.txt": ["LAMISIL 250", "LAMISIL 1% TOPICAL CREAM"],
    "43_duphaston_10.txt": ["DUPHASTON 10"],
    "44_folic_acid_5.txt": ["FOLIC ACID (EIPICO) 5 MG", "FOLIC ACID (EL NILE) 5 MG", "FOLIC ACID (AMOUN) 5MG", "FOLIC 0.8"],
    "45_ferrotron.txt": ["FERROTRON 30", "FERROTRON 20"],
    "46_milga.txt": ["MILGA 40", "MILGA ADVANCE"],
    "47_smecta.txt": ["SMECTA 20% SUSPENSION 120", "SMECTA"],
    "48_enterogermina.txt": ["ENTEROGERMINA 2 BILLION"],
    "49_strepsils.txt": ["STREPSILS HONEY", "STREPSILS COOL"],
    "50_comtrex.txt": ["COMTREX ACUTE HEAD COLD"],
}

BRAND_BY_FILE = {
    "01_panadol_paracetamol_500.txt": "Panadol Advance",
    "02_panadol_cold_flu_day.txt": "Panadol Cold & Flu Day",
    "03_panadol_sinus_relief_pe.txt": "Panadol Sinus Relief PE",
    "51_panadol_extra.txt": "Panadol Extra",
    "52_panadol_joint.txt": "Panadol Joint",
    "53_panadol_migraine.txt": "Panadol Migraine",
    "54_panadol_acute_head_cold.txt": "Panadol Acute Head Cold",
    "55_panadol_cold_flu_all_in_one.txt": "Panadol Cold & Flu All In One",
    "56_panadol_vapour_release.txt": "Panadol Cold & Flu Vapour Release",
    "04_cetal_cold_flu_day.txt": "Cetal Cold & Flu Day",
    "05_flurest_n.txt": "Flurest-N",
    "06_congestal.txt": "Congestal",
    "07_sine_up.txt": "Sine-Up",
    "08_bradozen.txt": "Bradozen",
    "09_notussil.txt": "Notussil",
    "10_tussigreen.txt": "Tussigreen",
    "11_daflon_500.txt": "Daflon 500",
    "12_daflon_1000.txt": "Daflon 1000",
    "13_controloc_20.txt": "Controloc 20",
    "14_jeparilon.txt": "Jeparilon",
    "15_anaseziago.txt": "Anaseziago",
    "16_adol_500.txt": "Adol",
    "17_cataflam_25.txt": "Cataflam",
    "18_brufen_400.txt": "Brufen",
    "19_antinal.txt": "Antinal",
    "20_flagyl_500.txt": "Flagyl",
    "21_augmentin_1g.txt": "Augmentin",
    "22_hibiotic_1g.txt": "Hi-Biotic",
    "23_zithrokan_500.txt": "Zithrokan",
    "24_ciprobay_500.txt": "Ciprobay",
    "25_risek_20.txt": "Risek",
    "26_gaviscon.txt": "Gaviscon Advance",
    "27_motilium_10.txt": "Motilium",
    "28_buscopan.txt": "Buscopan",
    "29_concor_5.txt": "Concor",
    "30_norvasc_5.txt": "Norvasc",
    "31_glucophage_500.txt": "Glucophage",
    "32_amaryl_2.txt": "Amaryl",
    "33_lipitor_20.txt": "Lipitor",
    "34_plavix_75.txt": "Plavix",
    "35_aspocid_75.txt": "Aspocid",
    "36_telfast_120.txt": "Telfast",
    "37_claritine_10.txt": "Claritine",
    "38_otrivin_0_1.txt": "Otrivin",
    "39_ventolin_inhaler.txt": "Ventolin",
    "40_singulair_10.txt": "Singulair",
    "41_diflucan_150.txt": "Diflucan",
    "42_lamisil_250.txt": "Lamisil",
    "43_duphaston_10.txt": "Duphaston",
    "44_folic_acid_5.txt": "Folic Acid",
    "45_ferrotron.txt": "Ferrotron",
    "46_milga.txt": "Milga",
    "47_smecta.txt": "Smecta",
    "48_enterogermina.txt": "Enterogermina",
    "49_strepsils.txt": "Strepsils",
    "50_comtrex.txt": "Comtrex",
}

# Canonical actives for sheet matching (INN-first). Sheets key by INN, not brand.
INN_BY_FILE: dict[str, list[str]] = {
    "01_panadol_paracetamol_500.txt": ["Paracetamol"],
    "02_panadol_cold_flu_day.txt": ["Paracetamol", "Caffeine", "Phenylephrine"],
    "03_panadol_sinus_relief_pe.txt": ["Paracetamol", "Phenylephrine"],
    "04_cetal_cold_flu_day.txt": ["Paracetamol", "Caffeine", "Phenylephrine"],
    "05_flurest_n.txt": ["Paracetamol", "Chlorpheniramine Maleate", "Phenylephrine"],
    "06_congestal.txt": ["Paracetamol", "Pseudoephedrine", "Chlorpheniramine Maleate"],
    "07_sine_up.txt": ["Pseudoephedrine", "Triprolidine"],
    "08_bradozen.txt": ["Bromelain", "Trypsin", "Rutoside"],
    "09_notussil.txt": ["Levodropropizine"],
    "10_tussigreen.txt": ["Ivy", "Ivy Extract"],
    "11_daflon_500.txt": ["Diosmin", "Hesperidin"],
    "12_daflon_1000.txt": ["Diosmin", "Hesperidin"],
    "13_controloc_20.txt": ["Pantoprazole"],
    # Egyptian JEPARILON catalog = famotidine + antacids (not topical heparinoid)
    "14_jeparilon.txt": ["Famotidine", "Calcium Carbonate", "Magnesium Hydroxide"],
    "15_anaseziago.txt": ["Lidocaine", "Prilocaine"],
    "16_adol_500.txt": ["Paracetamol"],
    "17_cataflam_25.txt": ["Diclofenac Potassium", "Diclofenac"],
    "18_brufen_400.txt": ["Ibuprofen"],
    "19_antinal.txt": ["Nifuroxazide"],
    "20_flagyl_500.txt": ["Metronidazole"],
    "21_augmentin_1g.txt": ["Amoxicillin", "Clavulanic Acid"],
    "22_hibiotic_1g.txt": ["Amoxicillin", "Clavulanic Acid"],
    "23_zithrokan_500.txt": ["Azithromycin"],
    "24_ciprobay_500.txt": ["Ciprofloxacin"],
    "25_risek_20.txt": ["Omeprazole"],
    "26_gaviscon.txt": ["Sodium Alginate", "Potassium Bicarbonate", "Alginate"],
    "27_motilium_10.txt": ["Domperidone"],
    "28_buscopan.txt": ["Hyoscine Butylbromide", "Hyoscine"],
    "29_concor_5.txt": ["Bisoprolol"],
    "30_norvasc_5.txt": ["Amlodipine"],
    "31_glucophage_500.txt": ["Metformin"],
    "32_amaryl_2.txt": ["Glimepiride"],
    "33_lipitor_20.txt": ["Atorvastatin"],
    "34_plavix_75.txt": ["Clopidogrel"],
    "35_aspocid_75.txt": ["Aspirin", "Acetylsalicylic Acid"],
    "36_telfast_120.txt": ["Fexofenadine"],
    "37_claritine_10.txt": ["Loratadine"],
    "38_otrivin_0_1.txt": ["Xylometazoline"],
    "39_ventolin_inhaler.txt": ["Salbutamol", "Albuterol"],
    "40_singulair_10.txt": ["Montelukast"],
    "41_diflucan_150.txt": ["Fluconazole"],
    "42_lamisil_250.txt": ["Terbinafine"],
    "43_duphaston_10.txt": ["Dydrogesterone"],
    "44_folic_acid_5.txt": ["Folic Acid"],
    "45_ferrotron.txt": ["Iron", "Ferrous", "Folic Acid"],
    "46_milga.txt": ["Benfotiamine", "Pyridoxine", "Cyanocobalamin"],
    "47_smecta.txt": ["Diosmectite"],
    "48_enterogermina.txt": ["Bacillus Clausii"],
    "49_strepsils.txt": ["Amylmetacresol", "Dichlorobenzyl"],
    "50_comtrex.txt": ["Paracetamol", "Pseudoephedrine", "Chlorpheniramine Maleate"],
    "51_panadol_extra.txt": ["Paracetamol", "Caffeine"],
    "52_panadol_joint.txt": ["Paracetamol"],
    "53_panadol_migraine.txt": ["Paracetamol", "Caffeine"],
    "54_panadol_acute_head_cold.txt": ["Paracetamol", "Guaifenesin", "Phenylephrine"],
    "55_panadol_cold_flu_all_in_one.txt": ["Paracetamol", "Guaifenesin", "Phenylephrine"],
    "56_panadol_vapour_release.txt": ["Paracetamol", "Guaifenesin", "Phenylephrine"],
}

_INN_ALIASES: dict[str, list[str]] = {
    "PARACETAMOL": ["PARACETAMOL (ACETAMINOPHEN)", "PARACETAMOL", "ACETAMINOPHEN"],
    "ACETYLSALICYLIC ACID": ["ASPIRIN", "ACETYLSALICYLIC ACID"],
    "ASPIRIN": ["ASPIRIN", "ACETYLSALICYLIC ACID"],
    "PHENYLEPHRINE": ["PHENYLEPHRINE", "PHENYLEPHRINE HCL", "PHENYLEPHRINE HYDROCHLORIDE"],
    "PSEUDOEPHEDRINE": ["PSEUDOEPHEDRINE", "PSEUDOEPHEDRINE HCL", "PSEUDOEPHEDRINE HYDROCHLORIDE"],
    "CHLORPHENIRAMINE MALEATE": ["CHLORPHENIRAMINE", "CHLORPHENIRAMINE MALEATE", "CHLORPHENAMINE"],
    "DICLOFENAC": ["DICLOFENAC POTASSIUM", "DICLOFENAC SODIUM", "DICLOFENAC"],
    "DICLOFENAC POTASSIUM": ["DICLOFENAC POTASSIUM", "DICLOFENAC"],
    "CLAVULANIC ACID": ["CLAVULANIC ACID", "CLAVULANATE", "POTASSIUM CLAVULANATE"],
    "HYOSCINE BUTYLBROMIDE": ["HYOSCINE BUTYLBROMIDE", "HYOSCINE", "BUTYLSCOPOLAMINE"],
    "BISOPROLOL": ["BISOPROLOL", "BISOPROLOL FUMARATE"],
    "AMLODIPINE": ["AMLODIPINE", "AMLODIPINE BESYLATE"],
    "METFORMIN": ["METFORMIN", "METFORMIN HCL", "METFORMIN HYDROCHLORIDE"],
    "ATORVASTATIN": ["ATORVASTATIN", "ATORVASTATIN CALCIUM"],
    "FEXOFENADINE": ["FEXOFENADINE", "FEXOFENADINE HCL"],
    "XYLOMETAZOLINE": ["XYLOMETAZOLINE", "XYLOMETAZOLINE HCL"],
    "SALBUTAMOL": ["SALBUTAMOL", "ALBUTEROL", "SALBUTAMOL SULFATE"],
    "MONTELUKAST": ["MONTELUKAST", "MONTELUKAST SODIUM"],
    "TERBINAFINE": ["TERBINAFINE", "TERBINAFINE HCL"],
    "IVY": ["IVY", "IVY EXTRACT", "HEDERA HELIX"],
    "RUTOSIDE": ["RUTOSIDE", "RUTOSIDE TRIHYDRATE", "RUTIN"],
    "SODIUM ALGINATE": ["SODIUM ALGINATE", "ALGINATE"],
    "DIOSMECTITE": ["DIOSMECTITE", "DIOCTAHEDRAL SMECTITE", "SMECTITE"],
    "BACILLUS CLAUSII": ["BACILLUS CLAUSII", "BACILLUS CLAUSII SPORES"],
    "IRON": ["IRON", "FERROUS", "FERROUS FUMARATE", "FERROUS SULFATE"],
    "FERROUS": ["FERROUS", "IRON", "FERROUS FUMARATE"],
}


def _clean(val) -> str:
    if val is None:
        return ""
    text = str(val).replace("\xa0", " ").strip()
    text = re.sub(r"[ \t]+", " ", text)
    return text


def _clip(text: str, n: int = 500) -> str:
    text = _clean(text)
    if len(text) <= n:
        return text
    return text[: n - 1].rstrip() + "…"


def expand_inn_keys(names: list[str]) -> list[str]:
    """Expand display INNs into sheet-key candidates (deduped, order preserved)."""
    out: list[str] = []
    seen: set[str] = set()
    for raw in names:
        base = _clean(raw).upper()
        if not base:
            continue
        candidates = [base]
        candidates.extend(_INN_ALIASES.get(base, []))
        token = base.split()[0]
        if token and token != base:
            candidates.append(token)
            candidates.extend(_INN_ALIASES.get(token, []))
        for cand in candidates:
            c = _clean(cand).upper()
            if c and c not in seen:
                seen.add(c)
                out.append(c)
    return out


def load_medicin() -> list[dict]:
    wb = openpyxl.load_workbook(EXCEL / "FINAL MEDICIN.xlsx", read_only=True, data_only=True)
    ws = wb["All-Data"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows)]
    out = []
    for row in rows:
        if not row or row[1] in (None, ""):
            continue
        item = {header[i]: row[i] if i < len(row) else None for i in range(len(header))}
        out.append(item)
    wb.close()
    return out


def load_all_material() -> dict[str, dict]:
    wb = openpyxl.load_workbook(EXCEL / "ALL MATERIAL 1.xlsx", read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    by_mat: dict[str, dict] = {}
    for row in rows:
        if not row or row[0] in (None, ""):
            continue
        item = {
            str(header[i]) if header[i] is not None else f"c{i}": row[i]
            for i in range(len(header))
        }
        key = _clean(item.get("New Material")).upper()
        if key and key not in by_mat:
            by_mat[key] = item
    wb.close()
    return by_mat


def load_disease_map() -> dict[str, str]:
    path = EXCEL / "FINAL MATERIAL VS DISEASE.xlsx"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = ws.iter_rows(values_only=True)
    next(rows)
    out: dict[str, str] = {}
    for row in rows:
        if not row or row[0] in (None, ""):
            continue
        out[_clean(row[0]).upper()] = _clean(row[1])
    wb.close()
    return out


def load_interactions(limit_per_api: int = 20) -> dict[str, list[tuple[str, str]]]:
    wb = openpyxl.load_workbook(
        EXCEL / "NEW DRUG INTERACTION.xlsx", read_only=True, data_only=True
    )
    ws = wb["Sheet1"]
    rows = ws.iter_rows(values_only=True)
    next(rows)
    out: dict[str, list[tuple[str, str]]] = {}
    for row in rows:
        if not row or row[1] in (None, "") or row[2] in (None, ""):
            continue
        api1 = _clean(row[1]).upper()
        api2 = _clean(row[2])
        note = _clip(_clean(row[3]), 320)
        bucket = out.setdefault(api1, [])
        if len(bucket) < limit_per_api:
            bucket.append((api2, note))
        api2u = api2.upper()
        bucket2 = out.setdefault(api2u, [])
        if len(bucket2) < limit_per_api:
            bucket2.append((_clean(row[1]), note))
    wb.close()
    return out


def pick_sku(med_rows: list[dict], prefers: list[str]) -> dict | None:
    upper_rows = [(_clean(r.get("med")).upper(), r) for r in med_rows]
    for pref in prefers:
        pref_u = pref.upper()
        for med_u, row in upper_rows:
            if pref_u in med_u:
                return row
    # fallback: first prefer token word
    token = prefers[0].split()[0].upper()
    for med_u, row in upper_rows:
        if token in med_u:
            return row
    return None


def primary_materials(material_blob: str) -> list[str]:
    parts = [p.strip() for p in _clean(material_blob).split("*") if p.strip()]
    return parts


def find_material_row(materials: dict[str, dict], names: list[str]) -> dict | None:
    for key in expand_inn_keys(names):
        if key in materials:
            return materials[key]
        for mk, row in materials.items():
            if key in mk or mk in key:
                return row
    return None


def lookup_disease(disease_map: dict[str, str], names: list[str]) -> str:
    chunks: list[str] = []
    seen: set[str] = set()
    for key in expand_inn_keys(names):
        hit = disease_map.get(key, "")
        if not hit:
            for dk, dv in disease_map.items():
                if key in dk or dk in key:
                    hit = dv
                    break
        if hit and hit not in seen:
            seen.add(hit)
            chunks.append(f"[{key}] {hit}")
        if len(chunks) >= 3:
            break
    return "\n".join(chunks)


def interaction_lines(
    inter_map: dict[str, list[tuple[str, str]]],
    names: list[str],
    *,
    limit: int = 12,
) -> list[str]:
    lines: list[str] = []
    seen: set[tuple[str, str]] = set()

    def _add(label: str, pairs: list[tuple[str, str]]) -> None:
        nonlocal lines
        for api2, note in pairs:
            sig = (label, api2.lower())
            if sig in seen:
                continue
            seen.add(sig)
            lines.append(f"- {label.title()} + {api2}: {note}")
            if len(lines) >= limit:
                return

    for key in expand_inn_keys(names):
        if len(lines) >= limit:
            break
        if key in inter_map:
            _add(key, inter_map[key])
            continue
        for mk, pairs in inter_map.items():
            if key in mk or mk.startswith(key) or key.startswith(mk):
                _add(mk, pairs)
                if len(lines) >= limit:
                    break
    return lines


def _resolve_lookup_names(file_name: str, sku: dict) -> tuple[list[str], str]:
    explicit = list(INN_BY_FILE.get(file_name) or [])
    catalog = primary_materials(_clean(sku.get("New Material")))
    merged: list[str] = []
    for name in explicit + catalog:
        if name and name not in merged:
            merged.append(name)
    display = " * ".join(explicit) if explicit else (" * ".join(catalog) if catalog else "-")
    return merged, display


def render_leaflet(
    *,
    brand: str,
    sku: dict,
    inn_display: str,
    material_row: dict | None,
    disease: str,
    interactions: list[str],
    note: str,
) -> str:
    med = _clean(sku.get("med"))
    catalog_material = _clean(sku.get("New Material")) or "-"
    uses = _clean(sku.get("Uses")) or "-"
    conc = _clean(sku.get("CONC.")) or "Not stated in catalog row"
    dose = _clean(sku.get("DOSE")) or "Not stated numerically in catalog row — follow pack/doctor"
    min_age = _clean(sku.get("Min. age")) or "See pack"
    form = _clean(sku.get("Dosage Form")) or "-"
    company = _clean(sku.get("Company")) or "-"
    price = _clean(sku.get("Price")) or "Not listed"
    cat1 = _clean(sku.get("1Category")) or "-"
    cat2 = _clean(sku.get("2Category")) or "-"
    cat3 = _clean(sku.get("3Category")) or "-"
    cross = _clean(sku.get("Cross-selling")) or "None listed"

    approved = _clip(material_row.get("Approved Uses") if material_row else "", 900) or "See labeled uses"
    off_label = _clip(material_row.get("Off-Label Uses") if material_row else "", 500) or "Not listed"
    black_box = _clip(material_row.get("BLACK BOX") if material_row else "", 700) or "None listed in material sheet"
    notes = _clip(material_row.get("NOTES") if material_row else "", 700) or "None listed"
    patient_en = _clip(material_row.get("PATIENT EDUCATION") if material_row else "", 700) or "Follow pharmacist counseling"
    se_maj = _clip(material_row.get("SIDE EFFECTS (major)") if material_row else "", 600) or "See pack"
    se_min = _clip(material_row.get("SIDE EFFECTS (minor)") if material_row else "", 600) or "See pack"
    disease_txt = _clip(
        disease or (material_row.get("INTERACTION DISEASE") if material_row else ""),
        1400,
    ) or "See pack contraindications"

    inter_block = "\n".join(interactions) if interactions else "- No top pairs extracted for this API in the interaction sheet sample."

    return f"""{med}
Patient Information Leaflet — Egyptian Commercial E2E Corpus (Excel-enriched)
Brand: {brand}
Catalog SKU: {med}
INN / actives: {inn_display}
Catalog New Material: {catalog_material}
Strength / form: CONC={conc} | Form={form} | Min.age={min_age}
Manufacturer / market: {company}
Price (catalog): {price}
Categories: {cat1} > {cat2} > {cat3}
Source files: FINAL MEDICIN.xlsx + ALL MATERIAL 1.xlsx + NEW DRUG INTERACTION.xlsx + FINAL MATERIAL VS DISEASE.xlsx
Corpus note: Enriched from local Egyptian pharmacy Excel sheets for AlgoRAG testing. Verify against official PIL/MoH label before clinical use.
Match note: {note}

1. What {brand} is used for
Catalog uses: {uses}
Approved uses (material sheet): {approved}
Off-label mentions (material sheet): {off_label}

2. Dosage — Adults
Catalog DOSE field: {dose}
Practical counseling: take exactly as prescribed or as on the Egyptian pack leaflet. Do not exceed labeled maxima. Complete antibiotic courses when prescribed.

3. Dosage — Children
Minimum age in catalog: {min_age}
Use paediatric SKUs/suspensions when required. Adult tablets/capsules are not automatically suitable for children.

4. Do not take {brand} if / disease cautions
From material/disease sheets:
{disease_txt}

5. Warnings and precautions
BLACK BOX / serious warnings (material sheet): {black_box}
Notes (material sheet): {notes}
Patient education excerpt: {patient_en}

6. Interactions (from NEW DRUG INTERACTION.xlsx)
{inter_block}
Also avoid duplicating the same actives in other cold/pain/antibiotic products.

7. Pregnancy
Follow obstetric advice. Many NSAIDs are avoided in the third trimester; live vaccines and some antifungals have special restrictions. If uncertain, ask a pharmacist/doctor and check the official PIL.

8. Breastfeeding
Check the official PIL. Prefer pharmacist confirmation for combination cold products, muscle relaxants, and strong antibiotics.

9. Supplements / cross-selling
Catalog cross-selling: {cross}
Supportive tips: separate minerals from quinolones/tetracyclines by 2–4 hours; probiotics may be timed away from antibiotics; ORS with antidiarrheals when dehydrated.

10. Possible side effects
Major (material sheet): {se_maj}
Minor (material sheet): {se_min}

11. Storage
Store as labeled on the Egyptian pack (commonly below 25–30°C, dry, away from children). Reconstituted suspensions may need refrigeration — check the bottle.
"""


def main() -> None:
    print("Loading Excel sheets...")
    med_rows = load_medicin()
    materials = load_all_material()
    disease_map = load_disease_map()
    inter_map = load_interactions()
    print(f"medicin={len(med_rows)} materials={len(materials)} disease={len(disease_map)} inter_keys={len(inter_map)}")

    OUT.mkdir(parents=True, exist_ok=True)
    index = [
        "# Egyptian commercial leaflets (Excel-enriched)",
        "",
        "Generated from local Excel sheets + leaflet templates.",
        "Clinical sheet matching uses INN / actives (not brand names).",
        "",
        "| File | Brand | INN | Matched SKU | Interactions |",
        "| --- | --- | --- | --- | --- |",
    ]
    missing = []
    for file_name, prefers in SKU_PREFERS.items():
        brand = BRAND_BY_FILE[file_name]
        sku = pick_sku(med_rows, prefers)
        if sku is None:
            missing.append(file_name)
            note = "No exact FINAL MEDICIN match — keep previous manual leaflet content"
            # skip overwrite if no match? still try token
            print("MISSING", file_name)
            continue
        lookup_names, inn_display = _resolve_lookup_names(file_name, sku)
        mat_row = find_material_row(materials, lookup_names)
        disease = lookup_disease(disease_map, lookup_names)
        inters = interaction_lines(inter_map, lookup_names, limit=12)
        note = f"Matched prefers={prefers[0]}; sheet lookup INNs={', '.join(lookup_names[:6])}"
        text = render_leaflet(
            brand=brand,
            sku=sku,
            inn_display=inn_display,
            material_row=mat_row,
            disease=disease,
            interactions=inters,
            note=note,
        )
        (OUT / file_name).write_text(text, encoding="utf-8")
        index.append(
            f"| `{file_name}` | {brand} | {inn_display} | {_clean(sku.get('med'))} | {len(inters)} |"
        )
        print(
            "OK",
            file_name,
            "->",
            _clean(sku.get("med")),
            f"inter={len(inters)} mat={bool(mat_row)} dis={bool(disease)}",
        )
    # Tussigreen: no exact brand; enrich via closest cough SKU + Ivy INN
    sku = pick_sku(med_rows, ["TUSSIGLOBE", "TUSSIPECT-N", "TUSSISTOP"])
    if sku is not None:
        brand = "Tussigreen"
        lookup_names, inn_display = _resolve_lookup_names("10_tussigreen.txt", sku)
        mat_row = find_material_row(materials, lookup_names)
        disease = lookup_disease(disease_map, lookup_names)
        inters = interaction_lines(inter_map, lookup_names, limit=12)
        text = render_leaflet(
            brand=brand,
            sku=sku,
            inn_display=inn_display or "Ivy Leaf Extract (Hedera helix)",
            material_row=mat_row,
            disease=disease,
            interactions=inters,
            note=(
                "No TUSSIGREEN SKU in FINAL MEDICIN; enriched using closest cough "
                "catalog SKU + Ivy INN sheet lookup"
            ),
        )
        text = text.replace(
            _clean(sku.get("med")),
            f"TUSSIGREEN (corpus proxy via {_clean(sku.get('med'))})",
            1,
        )
        (OUT / "10_tussigreen.txt").write_text(text, encoding="utf-8")
        index = [ln for ln in index if "10_tussigreen" not in ln]
        index.append(
            f"| `10_tussigreen.txt` | Tussigreen | {inn_display} | proxy:{_clean(sku.get('med'))} | {len(inters)} |"
        )
        print("OK proxy tussigreen", _clean(sku.get("med")), f"inter={len(inters)}")

    (OUT / "README.md").write_text("\n".join(index) + "\n", encoding="utf-8")
    print("Done. missing exact:", missing)


if __name__ == "__main__":
    main()
